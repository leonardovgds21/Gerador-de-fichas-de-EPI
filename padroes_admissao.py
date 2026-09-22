"""
padroes_admissao.py
--------------------
Módulo responsável por:
    1. Listar os "padrões de admissão" disponíveis: um por arquivo .xlsx
       dentro da pasta modelos_x_cliente (pasta raiz do projeto), usando o
       nome do arquivo (sem extensão) como nome do padrão.
    2. Ler, de um desses arquivos, a lista de nomes de equipamentos (coluna
       E da planilha), para popular automaticamente a lista de itens da
       interface ao selecionar um cliente.

Todas as comparações de texto feitas aqui (nome do padrão x nome da aba,
nome do equipamento x base equipamentos.xlsx) ignoram maiúsculas/minúsculas
e acentuação, conforme solicitado.
"""

import os
import re
import unicodedata
import openpyxl

from caminhos import obter_pasta_base

# Pasta base do sistema: pasta do projeto ao rodar como script, ou pasta do
# .exe ao rodar como executável gerado pelo PyInstaller (veja caminhos.py).
PASTA_DO_PROJETO = obter_pasta_base()
PASTA_MODELOS_CLIENTE = os.path.join(PASTA_DO_PROJETO, "modelos_x_cliente")

# Texto do cabeçalho da coluna de itens nas planilhas de cliente. É usado
# para localizar automaticamente em qual linha a lista de equipamentos
# começa (essa linha varia de arquivo para arquivo).
TEXTO_CABECALHO_ITENS = "Relação de Uniformes/EPIs"

# Colunas fixas dentro das planilhas de cliente (1 = A, 2 = B, ...).
COLUNA_QUANTIDADE = 3  # coluna C ("QTD.")
COLUNA_ITEM = 5         # coluna E ("Relação de Uniformes/EPIs")


def normalizar(texto) -> str:
    """
    Remove acentuação e diferenças de maiúsculas/minúsculas de um texto, para
    permitir comparações tolerantes (ex.: "ADM Campo" == "adm campo" ==
    "ADM CAMPO"). Usada tanto para comparar nomes de equipamentos quanto
    nomes de padrões/abas.
    """
    texto = "" if texto is None else str(texto)
    texto = texto.strip()
    sem_acentos = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    return sem_acentos.casefold()


def _normalizar_chave(texto) -> str:
    """Normalização mais agressiva (remove também espaços, underscores e
    pontuação), usada só para comparar nome de arquivo com nome de aba, já
    que essas duas grafias costumam variar mais entre si (ex.: "KWS_GDM"
    no arquivo e "KWS _ GDM" na aba)."""
    return re.sub(r"[^a-z0-9]", "", normalizar(texto))


def listar_padroes_admissao() -> list:
    """
    Retorna, em ordem alfabética, os nomes dos padrões de admissão
    disponíveis — um por arquivo .xlsx dentro de modelos_x_cliente — usando
    o nome do arquivo sem a extensão (mesmo nome mostrado na interface).
    """
    if not os.path.isdir(PASTA_MODELOS_CLIENTE):
        return []

    nomes = [
        os.path.splitext(nome_arquivo)[0]
        for nome_arquivo in os.listdir(PASTA_MODELOS_CLIENTE)
        # ignora arquivos temporários de lock do Excel (ex.: "~$Cargill.xlsx")
        if nome_arquivo.lower().endswith(".xlsx") and not nome_arquivo.startswith("~$")
    ]
    return sorted(nomes, key=normalizar)


def _selecionar_aba(planilha, nome_padrao: str):
    """
    Algumas planilhas de cliente têm mais de uma aba (ex.: uma aba genérica
    de "ENGENHARIA (OBRA)" compartilhada entre arquivos, e outra específica
    do site). Esta função escolhe a aba cujo nome corresponde ao padrão
    selecionado (comparação tolerante a acentos/caixa/espaços); se nenhuma
    aba corresponder, usa a primeira aba do arquivo.
    """
    alvo = _normalizar_chave(nome_padrao)
    for nome_aba in planilha.sheetnames:
        if _normalizar_chave(nome_aba) == alvo:
            return planilha[nome_aba]
    return planilha[planilha.sheetnames[0]]


def _valor_e_numero_ou_vazio(valor) -> bool:
    """
    Indica se um valor da coluna "QTD." parece uma quantidade (número ou
    vazio). É usado para diferenciar linhas de itens de EPI das linhas de
    observações finais da planilha (ex.: "TICKET LOG", "CLIENTE", "NOME:"),
    que trazem texto nessa coluna em vez de número.
    """
    if valor is None:
        return True
    if isinstance(valor, (int, float)):
        return True
    texto = str(valor).strip()
    if texto == "":
        return True
    try:
        float(texto.replace(",", "."))
        return True
    except ValueError:
        return False


def carregar_itens_do_padrao(nome_padrao: str) -> list:
    """
    Lê modelos_x_cliente/<nome_padrao>.xlsx e retorna a lista de nomes de
    equipamentos encontrados na coluna E, na ordem em que aparecem na
    planilha. A leitura é interrompida assim que uma linha de observação
    final é encontrada (identificada por trazer texto, em vez de número ou
    vazio, na coluna "QTD.").
    """
    caminho_arquivo = os.path.join(PASTA_MODELOS_CLIENTE, f"{nome_padrao}.xlsx")
    if not os.path.exists(caminho_arquivo):
        raise FileNotFoundError(f"Padrão de admissão não encontrado: {caminho_arquivo}")

    planilha = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    aba = _selecionar_aba(planilha, nome_padrao)

    # Localiza a linha do cabeçalho de itens ("Relação de Uniformes/EPIs"),
    # já que ela não fica sempre na mesma linha em todos os arquivos.
    linha_cabecalho = None
    for linha in aba.iter_rows(min_row=1, max_row=10):
        valor_coluna_item = linha[COLUNA_ITEM - 1].value
        if normalizar(valor_coluna_item) == normalizar(TEXTO_CABECALHO_ITENS):
            linha_cabecalho = linha[0].row
            break

    if linha_cabecalho is None:
        raise ValueError(
            f'Não foi possível localizar o cabeçalho "{TEXTO_CABECALHO_ITENS}" '
            f'na planilha do padrão "{nome_padrao}".'
        )

    nomes_equipamentos = []
    for linha in aba.iter_rows(min_row=linha_cabecalho + 1, max_row=aba.max_row):
        valor_quantidade = linha[COLUNA_QUANTIDADE - 1].value
        valor_item = linha[COLUNA_ITEM - 1].value

        # Linha de observação final (ex.: "TICKET LOG", "CLIENTE", "NOME:")
        # -> a lista de equipamentos sempre termina antes dessas linhas.
        if not _valor_e_numero_ou_vazio(valor_quantidade):
            break

        if valor_item not in (None, ""):
            nomes_equipamentos.append(str(valor_item).strip())

    return nomes_equipamentos
